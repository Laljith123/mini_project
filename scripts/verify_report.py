import docx

doc = docx.Document('CPPE_Mini_Project_Report_Laljith_V.docx')
print('=== SUMMARY OF PARAGRAPH CHECKS ===')
sizes = set()
alignments = set()
bolds = set()

for i, p in enumerate(doc.paragraphs):
    alignments.add(str(p.alignment))
    for r in p.runs:
        if r.font.size:
            sizes.add(r.font.size.pt)
        bolds.add(r.bold)

print('Font sizes used in document:', sorted(list(sizes)))
print('Paragraph alignments used:', alignments)
print('Bold states used:', bolds)

# Verify key phrases in document
full_text = '\n'.join([p.text for p in doc.paragraphs])
for phrase in [
    'AI-POWERED RESUME AND JOB DESCRIPTION GENERATOR',
    'LALJITH V',
    '23UCB031',
    'Dr. M. AUXILIA',
    'DEPARTMENT OF COMPUTER SCIENCE AND BUSINESS SYSTEMS',
    'SRI MANAKULA VINAYAGAR ENGINEERING COLLEGE',
    'BONAFIDE CERTIFICATE',
    'ACKNOWLEDGEMENT',
    'TABLE OF CONTENTS',
    'DATA ENGINEERING',
    'CLOUD ENGINEERING',
    'AI ENGINEERING',
    'SECURITY ENGINEERING',
    '15 . 10 . 2026'
]:
    present = phrase in full_text
    print(f'Verification: "{phrase}" -> {"FOUND" if present else "MISSING!"}')

# Also check tables text
tbl_text = ''
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            tbl_text += cell.text + ' '

for phrase in ['LALJITH V', 'REGISTER NO: 23UCB031', 'PROJECT COORDINATOR', 'HEAD OF THE DEPARTMENT']:
    present = phrase in tbl_text
    print(f'Table Verification: "{phrase}" -> {"FOUND" if present else "MISSING!"}')
